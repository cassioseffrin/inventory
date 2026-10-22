#!/usr/bin/env python3
"""
Merge da planilha de inventário (aba COLABORADORES, formato "card") no banco `inventario`.

Cada card tem: COLABORADOR / CPF / SETOR / GERENTE / (Nº PATRIMÔNIO | ITEM | Nº MEC)... / RAMAL / E-MAIL
[/ CONTATO CELULAR]. Os cards ficam lado a lado (4 por linha) e são localizados pela célula "COLABORADOR".

Regras (o script é idempotente — pode ser reexecutado):
  Funcionário
    - chave = CPF normalizado (11 dígitos). Existe → update (nome, setor, e-mail, ramal, celular;
      valores vazios na planilha não apagam o banco). Não existe → insert com matrícula Fnnnn.
    - card sem CPF: tenta casar pelo e-mail; cards "ESTOQUE ..." viram itens em estoque sem responsável.
    - mesma pessoa em mais de um card: vale o primeiro card.
    - GERENTE: casado com um funcionário já cadastrado pelo nome completo ou pelo primeiro nome
      (precisa ser único); grava funcionarios.gerente_id. Sem match → apenas reportado.
  Item (nº de patrimônio e MAC são únicos no banco; em conflitos na planilha vale a primeira ocorrência)
    - MAC normalizado (aa:bb:cc:dd:ee:ff). MAC inválido é ignorado e guardado nas observações.
    - casa primeiro pelo MAC (itens.mac / itens_interfaces.mac), depois pelo nº de patrimônio.
    - existe → update (responsável, setor, tipo, status; patrimônio só substitui o provisório;
      MAC e modelo só preenchem quando vazios). Não existe → insert (sem patrimônio → PROV-nnnnnn).
    - itens sem patrimônio E sem MAC (mouse/fone "NÃO TEM") são ignorados, a menos que se use
      --incluir-sem-identificacao (nesse caso NÃO é idempotente: cada execução cria de novo).
  Todas as alterações vão para itens_historico com usuario = 'planilha'.

Uso:
  python3 -m venv .venv && .venv/bin/pip install -r scripts/merge_planilha/requirements.txt
  .venv/bin/python scripts/merge_planilha/merge_colaboradores.py data_merge/inventory17.10.26.xlsx            # simulação
  .venv/bin/python scripts/merge_planilha/merge_colaboradores.py data_merge/inventory17.10.26.xlsx --aplicar  # grava

A conexão vem de INV_DATABASE_URL (variável de ambiente ou .env na raiz do projeto).
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import psycopg2
import psycopg2.extras

USUARIO = 'planilha'
ABA = 'COLABORADORES'

# MACs de adaptadores virtuais que aparecem em várias máquinas (02:00:4c:4f:4f:50 = Microsoft Loopback "LOOP")
MACS_IGNORADOS = {'02:00:4c:4f:4f:50', '00:00:00:00:00:00', 'ff:ff:ff:ff:ff:ff'}

TIPO_POR_ITEM = {
    'NOTEBOOK': 'notebook',
    'CPU': 'desktop',
    'IMPRESSORA': 'impressora',
    'MAQ PAGAMENTO': 'smart_pos',
}


# ---------------------------------------------------------------- normalização
def ascii_upper(v) -> str:
    if v is None:
        return ''
    s = str(v).strip()
    return unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().upper()


def texto(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r'\s+', ' ', str(v)).strip()
    return s or None


def so_digitos(v) -> str:
    return re.sub(r'\D', '', str(v or ''))


def normalizar_cpf(v) -> str | None:
    d = so_digitos(v)
    return d if len(d) == 11 and not re.fullmatch(r'(\d)\1{10}', d) else None


def cpf_valido(d: str) -> bool:
    def dv(n):
        soma = sum(int(d[i]) * (n + 1 - i) for i in range(n))
        r = (soma * 10) % 11
        return 0 if r == 10 else r
    return dv(9) == int(d[9]) and dv(10) == int(d[10])


def normalizar_mac(v) -> str | None:
    """Como src/server/util.ts, mas recusa letras fora de a-f (ex.: '00-e0-4c-93-se-se') em vez de descartá-las."""
    hexa = re.sub(r'[\s:;.\-]', '', str(v or '')).lower()
    if not re.fullmatch(r'[0-9a-f]{12}', hexa):
        return None
    mac = ':'.join(hexa[i:i + 2] for i in range(0, 12, 2))
    return None if mac in MACS_IGNORADOS else mac


def normalizar_patrimonio(v) -> str | None:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = texto(v)
    return s if s and re.fullmatch(r'\d+', s) else None


def mascara_telefone(v) -> str | None:
    """Mesma regra de src/lib/telefone.ts; com vários números, usa o primeiro."""
    s = texto(v)
    if not s:
        return None
    partes = [p for p in re.split(r'(?=\()|/|;|,', s) if so_digitos(p)]
    d = so_digitos(partes[0] if partes else s).lstrip('0')[:11]
    if len(d) < 10:
        return None
    corte = 5 if len(d) == 11 else 4
    return f'({d[:2]}) {d[2:2 + corte]}-{d[2 + corte:]}'


def nome_proprio(v) -> str | None:
    s = texto(v)
    if not s:
        return None
    minusc = {'da', 'de', 'do', 'das', 'dos', 'e'}
    return ' '.join(p.lower() if i and p.lower() in minusc else p.capitalize() for i, p in enumerate(s.split(' ')))


# ---------------------------------------------------------------- leitura dos cards
def ler_cards(caminho: str) -> list[dict]:
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[ABA]
    maxr, maxc = ws.max_row, ws.max_column

    def V(r, c):
        return ws.cell(r, c).value if 1 <= r <= maxr and 1 <= c <= maxc else None

    cards = []
    for r in range(1, maxr + 1):
        for c in range(1, maxc + 1):
            if ascii_upper(V(r, c)) != 'COLABORADOR':
                continue
            card = {
                'ref': f'{ABA}!{openpyxl.utils.get_column_letter(c)}{r}',
                'nome': texto(V(r, c + 1)),
                'cpf_bruto': None, 'setor': None, 'gerente': None,
                'ramal': None, 'email': None, 'celular': None, 'itens': [],
            }
            modo = 'cabecalho'
            rr = r + 1
            while rr <= maxr:
                lab, lab2 = ascii_upper(V(rr, c)), ascii_upper(V(rr, c + 1))
                if lab == 'COLABORADOR':
                    break
                linha = [V(rr, c), V(rr, c + 1), V(rr, c + 2)]
                vazia = all(texto(x) is None for x in linha)
                if modo == 'cabecalho':
                    if lab == 'CPF':
                        card['cpf_bruto'] = V(rr, c + 1)
                    elif lab == 'SETOR':
                        card['setor'] = texto(V(rr, c + 1))
                    elif lab == 'GERENTE':
                        card['gerente'] = texto(V(rr, c + 1))
                    elif lab.startswith('N') and 'PATRIM' in lab:
                        modo = 'itens'
                elif modo == 'itens':
                    if lab == 'RAMAL' or lab2 in ('E-MAIL', 'EMAIL'):
                        modo = 'ramal'
                    elif vazia:
                        break
                    elif texto(linha[1]):
                        card['itens'].append({
                            'ref': f'{ABA}!{openpyxl.utils.get_column_letter(c)}{rr}',
                            'patrimonio_bruto': texto(linha[0]),
                            'item': texto(linha[1]).upper(),
                            'mac_bruto': texto(linha[2]),
                        })
                elif modo == 'ramal':
                    card['ramal'] = so_digitos(V(rr, c)) or None
                    card['email'] = (texto(V(rr, c + 1)) or '').lower() or None
                    modo = 'depois'
                elif modo == 'depois':
                    if 'CELULAR' in lab2 or 'CONTATO' in lab2:
                        card['celular'] = V(rr + 1, c + 1)
                        rr += 1
                    elif vazia:
                        break
                rr += 1
            cards.append(card)
    return cards


# ---------------------------------------------------------------- banco
def database_url() -> str:
    url = os.environ.get('INV_DATABASE_URL')
    if url:
        return url
    env = Path(__file__).resolve().parents[2] / '.env'
    if env.exists():
        for linha in env.read_text().splitlines():
            m = re.match(r'\s*INV_DATABASE_URL\s*=\s*(.*)\s*$', linha)
            if m:
                return m.group(1).strip().strip('"\'')
    sys.exit('INV_DATABASE_URL não definida (ambiente ou .env).')


class Merge:
    def __init__(self, cur, incluir_sem_id: bool):
        self.cur = cur
        self.incluir_sem_id = incluir_sem_id
        self.rel = {k: [] for k in (
            'func_criados', 'func_atualizados', 'gerentes', 'itens_criados', 'itens_atualizados',
            'itens_ignorados', 'conflitos', 'avisos')}
        self.patrimonios_vistos: dict[str, str] = {}
        self.macs_vistos: dict[str, str] = {}
        self.cpfs_vistos: dict[str, int] = {}
        self.gerente_do_card: dict[int, tuple[str, str]] = {}  # funcionario_id -> (gerente na planilha, ref)

    # -- util
    def one(self, sql, args=()):
        self.cur.execute(sql, args)
        return self.cur.fetchone()

    def historico(self, item, acao, mudancas):
        for campo, antigo, novo in mudancas or [(None, None, None)]:
            if campo == 'funcionario_id':
                antigo, novo = self.nome_func(antigo), self.nome_func(novo)
                acao_m = 'troca_responsavel' if acao == 'alteracao' else acao
            else:
                acao_m = acao
            self.cur.execute(
                'INSERT INTO inventario.itens_historico (item_id, numero_patrimonio, acao, campo, valor_antigo, valor_novo, usuario) '
                'VALUES (%s,%s,%s,%s,%s,%s,%s)',
                (item['id'], item['numero_patrimonio'], acao_m, campo,
                 None if antigo is None else str(antigo), None if novo is None else str(novo), USUARIO))

    def nome_func(self, fid):
        if fid is None:
            return None
        f = self.one('SELECT nome, matricula FROM inventario.funcionarios WHERE id=%s', (fid,))
        return f"{f['nome']} ({f['matricula']})" if f else f'#{fid}'

    # -- funcionários
    def proxima_matricula(self) -> str:
        r = self.one(r"SELECT coalesce(max(substring(matricula from '^F(\d+)$')::int), 0) + 1 AS n "
                     r"FROM inventario.funcionarios WHERE matricula ~ '^F\d+$'")
        return f"F{r['n']:04d}"

    def funcionario(self, card, cpf_por_email) -> tuple[int | None, bool]:
        """Retorna (funcionario_id, eh_estoque)."""
        nome_up = ascii_upper(card['nome'])
        if 'ESTOQUE' in nome_up:
            return None, True
        cpf = normalizar_cpf(card['cpf_bruto'])
        if not cpf and card['email']:
            cpf = cpf_por_email.get(card['email'])
            if not cpf:
                f = self.one('SELECT id FROM inventario.funcionarios WHERE lower(email)=%s', (card['email'],))
                if f:
                    self.rel['avisos'].append(f"{card['ref']} {card['nome']}: sem CPF, associado pelo e-mail {card['email']}")
                    return f['id'], False
            else:
                self.rel['avisos'].append(f"{card['ref']} {card['nome']}: sem CPF, usado o CPF de outro card com o mesmo e-mail")
        if not cpf:
            self.rel['avisos'].append(
                f"{card['ref']} {card['nome']!r}: CPF ausente/inválido ({card['cpf_bruto']!r}) — itens ficam sem responsável")
            return None, False
        if not cpf_valido(cpf):
            self.rel['avisos'].append(f"{card['ref']} {card['nome']}: CPF {cpf} com dígito verificador inválido (gravado assim mesmo)")

        dados = {
            'nome': nome_proprio(card['nome']),
            'setor': card['setor'],
            'email': card['email'],
            'ramal': card['ramal'],
            'celular': mascara_telefone(card['celular']),
        }
        # mesma pessoa em mais de um card: os dados do primeiro card valem
        if cpf in self.cpfs_vistos:
            return self.cpfs_vistos[cpf], False
        atual = self.one('SELECT * FROM inventario.funcionarios WHERE cpf=%s', (cpf,))
        if atual:
            self.cpfs_vistos[cpf] = atual['id']
            if card['gerente']:
                self.gerente_do_card[atual['id']] = (card['gerente'], card['ref'])
            patch = {k: v for k, v in dados.items() if v and (atual[k] or '') != v}
            if patch:
                sets = ', '.join(f'{k}=%s' for k in patch)
                self.cur.execute(f'UPDATE inventario.funcionarios SET {sets}, updated_at=now() WHERE id=%s',
                                 (*patch.values(), atual['id']))
                self.rel['func_atualizados'].append(
                    f"{atual['nome']} ({cpf}): " + ', '.join(f'{k}: {atual[k]!r} → {v!r}' for k, v in patch.items()))
            return atual['id'], False

        matricula = self.proxima_matricula()
        novo = self.one(
            'INSERT INTO inventario.funcionarios (nome, matricula, email, setor, cpf, ramal, celular) '
            'VALUES (%(nome)s, %(matricula)s, %(email)s, %(setor)s, %(cpf)s, %(ramal)s, %(celular)s) RETURNING id',
            {**dados, 'matricula': matricula, 'cpf': cpf})
        self.cpfs_vistos[cpf] = novo['id']
        if card['gerente']:
            self.gerente_do_card[novo['id']] = (card['gerente'], card['ref'])
        self.rel['func_criados'].append(f"{dados['nome']} ({cpf}) matrícula {matricula}")
        return novo['id'], False

    # -- gerentes (depois de todos os funcionários existirem)
    def gerentes(self):
        self.cur.execute('SELECT id, nome, gerente_id FROM inventario.funcionarios WHERE ativo')
        todos = self.cur.fetchall()
        chave = {f['id']: ascii_upper(f['nome']) for f in todos}
        gerente_atual = {f['id']: f['gerente_id'] for f in todos}

        def resolver(nome):
            n = ascii_upper(nome)
            exatos = [i for i, k in chave.items() if k == n]
            if exatos:
                return exatos
            return [i for i, k in chave.items() if k.split(' ')[0] == n.split(' ')[0] and k.startswith(n)]

        def ciclo(fid, gid):
            vistos = set()
            while gid and gid not in vistos:
                if gid == fid:
                    return True
                vistos.add(gid)
                gid = gerente_atual.get(gid)
            return False

        for fid, (nome_ger, ref) in self.gerente_do_card.items():
            achados = resolver(nome_ger)
            if len(achados) != 1:
                motivo = 'nenhum funcionário com esse nome' if not achados else f'{len(achados)} funcionários possíveis'
                self.rel['avisos'].append(f'{ref} {chave.get(fid)}: gerente {nome_ger!r} não associado ({motivo})')
                continue
            gid = achados[0]
            if gid == fid:
                continue  # o próprio gerente aparece com ele mesmo como gerente
            if gerente_atual.get(fid) == gid:
                continue
            if ciclo(fid, gid):
                self.rel['conflitos'].append(f'{ref} {chave.get(fid)}: gerente {nome_ger!r} geraria ciclo na hierarquia — ignorado')
                continue
            self.cur.execute('UPDATE inventario.funcionarios SET gerente_id=%s, updated_at=now() WHERE id=%s', (gid, fid))
            self.rel['gerentes'].append(f'{chave.get(fid)} → {chave[gid]}')
            gerente_atual[fid] = gid

    # -- itens
    def buscar_por_mac(self, mac):
        self.cur.execute(
            'SELECT * FROM inventario.itens i WHERE i.mac=%s '
            'OR EXISTS (SELECT 1 FROM inventario.itens_interfaces f WHERE f.item_id=i.id AND f.mac=%s) ORDER BY i.id',
            (mac, mac))
        return self.cur.fetchall()

    def item(self, card, it, func_id, estoque):
        pat = normalizar_patrimonio(it['patrimonio_bruto'])
        mac = normalizar_mac(it['mac_bruto'])
        mac_bruto = it['mac_bruto'] if it['mac_bruto'] and not ascii_upper(it['mac_bruto']).startswith('NAO') else None
        desc = f"{it['ref']} {card['nome']} / {it['item']} pat={pat or '-'} mac={mac or '-'}"
        obs_mac = f'MAC informado na planilha (inválido/virtual): {mac_bruto}' if mac_bruto and not mac else None
        if obs_mac:
            self.rel['avisos'].append(f'{desc}: MAC inválido/virtual {mac_bruto!r} ignorado')

        if not pat and not mac and not self.incluir_sem_id:
            self.rel['itens_ignorados'].append(desc)
            return
        # mesmo patrimônio/MAC repetido na planilha: vale a primeira ocorrência
        if mac and mac in self.macs_vistos:
            self.rel['conflitos'].append(f'{desc}: MAC já aparece em {self.macs_vistos[mac]} — linha ignorada')
            return
        if pat and pat in self.patrimonios_vistos:
            if not mac:
                self.rel['conflitos'].append(f'{desc}: patrimônio já aparece em {self.patrimonios_vistos[pat]} — linha ignorada')
                return
            self.rel['conflitos'].append(
                f'{desc}: patrimônio já aparece em {self.patrimonios_vistos[pat]} — item tratado só pelo MAC, sem esse patrimônio')
            pat = None
        if pat:
            self.patrimonios_vistos[pat] = desc
        if mac:
            self.macs_vistos[mac] = desc

        tipo = TIPO_POR_ITEM.get(ascii_upper(it['item']), 'outro')
        status = 'estoque' if estoque or not func_id else 'em_uso'

        atual = None
        if mac:
            achados = self.buscar_por_mac(mac)
            if len(achados) > 1:
                self.rel['avisos'].append(
                    f"{desc}: MAC em {len(achados)} itens ({', '.join(a['numero_patrimonio'] for a in achados)}) — usado o primeiro")
            atual = achados[0] if achados else None
        por_pat = self.one('SELECT * FROM inventario.itens WHERE numero_patrimonio=%s', (pat,)) if pat else None
        if atual and por_pat and por_pat['id'] != atual['id']:
            self.rel['conflitos'].append(
                f"{desc}: MAC está no item {atual['numero_patrimonio']} mas o patrimônio {pat} já é outro item (id {por_pat['id']}) "
                '— atualizado o item do MAC, patrimônio não alterado')
            pat_novo = None
        else:
            atual = atual or por_pat
            pat_novo = pat

        if atual:
            self.atualizar_item(atual, desc, pat_novo, mac, tipo, status, func_id, card['setor'], it['item'], obs_mac)
        else:
            self.criar_item(desc, pat, mac, tipo, status, func_id, card['setor'], it['item'], obs_mac, it['ref'])

    def atualizar_item(self, atual, desc, pat, mac, tipo, status, func_id, setor, item, obs_mac):
        patch = {}
        if pat and atual['numero_patrimonio'] != pat:
            if atual['patrimonio_provisorio']:
                patch['numero_patrimonio'] = pat
                patch['patrimonio_provisorio'] = False
            else:
                self.rel['conflitos'].append(
                    f"{desc}: item já tem patrimônio {atual['numero_patrimonio']} (planilha diz {pat}) — mantido")
        if mac and atual['mac'] != mac:
            if not atual['mac']:
                patch['mac'] = mac
            elif atual['numero_patrimonio'] == pat:
                self.rel['conflitos'].append(f"{desc}: item tem MAC {atual['mac']} (planilha diz {mac}) — mantido")
        if atual['tipo'] != tipo:
            if tipo == 'outro':
                # a planilha só diz "TELA/TECLADO..."; não rebaixa um tipo específico (ex.: vindo da patrimonial)
                self.rel['avisos'].append(f"{desc}: item é {atual['tipo']!r} no banco ({atual['descricao'] or atual['modelo']}) — tipo mantido")
            else:
                patch['tipo'] = tipo
        if func_id and atual['funcionario_id'] != func_id:
            patch['funcionario_id'] = func_id
        if setor and atual['setor'] != setor:
            patch['setor'] = setor
        if atual['status'] in ('em_uso', 'estoque') and atual['status'] != status:
            patch['status'] = status
        if not atual['modelo']:
            patch['modelo'] = item
        if obs_mac and obs_mac not in (atual['observacoes'] or ''):
            patch['observacoes'] = '\n'.join(filter(None, [atual['observacoes'], obs_mac]))
        if not patch:
            return
        sets = ', '.join(f'{k}=%s' for k in patch)
        self.cur.execute(f'UPDATE inventario.itens SET {sets}, updated_at=now() WHERE id=%s', (*patch.values(), atual['id']))
        item_ref = {'id': atual['id'], 'numero_patrimonio': patch.get('numero_patrimonio', atual['numero_patrimonio'])}
        self.historico(item_ref, 'alteracao', [(k, atual[k], v) for k, v in patch.items()])
        self.rel['itens_atualizados'].append(
            f"{desc} [{atual['numero_patrimonio']}]: " + ', '.join(f'{k}: {atual[k]!r} → {v!r}' for k, v in patch.items()))

    def criar_item(self, desc, pat, mac, tipo, status, func_id, setor, item, obs_mac, ref):
        provisorio = not pat
        if provisorio:
            n = self.one("SELECT nextval('inventario.patrimonio_provisorio_seq')::text AS n")['n']
            pat = f'PROV-{n.zfill(6)}'
        novo = self.one(
            'INSERT INTO inventario.itens (numero_patrimonio, patrimonio_provisorio, mac, tipo, status, funcionario_id, setor, modelo, observacoes) '
            'VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id, numero_patrimonio',
            (pat, provisorio, mac, tipo, status, func_id, setor, item, obs_mac))
        self.historico(novo, 'importacao', [('origem', None, f'planilha {ref}')])
        self.rel['itens_criados'].append(f'{desc} → {pat}')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('planilha')
    ap.add_argument('--aplicar', action='store_true', help='grava no banco (sem isso, só simula e faz rollback)')
    ap.add_argument('--incluir-sem-identificacao', action='store_true',
                    help='cria também itens sem patrimônio e sem MAC (não idempotente)')
    ap.add_argument('--detalhe', action='store_true', help='lista cada item criado/atualizado/ignorado')
    a = ap.parse_args()

    cards = ler_cards(a.planilha)
    cpf_por_email = {}
    for c in cards:
        cpf = normalizar_cpf(c['cpf_bruto'])
        if cpf and c['email']:
            cpf_por_email.setdefault(c['email'], cpf)

    con = psycopg2.connect(database_url(), application_name='merge-planilha')
    try:
        with con:
            cur = con.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
            m = Merge(cur, a.incluir_sem_identificacao)
            for card in cards:
                func_id, estoque = m.funcionario(card, cpf_por_email)
                for it in card['itens']:
                    m.item(card, it, func_id, estoque)
            m.gerentes()
            if not a.aplicar:
                con.rollback()
    finally:
        con.close()

    print(f"{'APLICADO' if a.aplicar else 'SIMULAÇÃO (use --aplicar para gravar)'} — {len(cards)} cards lidos de {a.planilha}\n")
    sempre = ('conflitos', 'avisos', 'func_criados', 'func_atualizados')
    for k, linhas in m.rel.items():
        print(f'== {k}: {len(linhas)}')
        if a.detalhe or k in sempre:
            for l in linhas:
                print('  ', l)


if __name__ == '__main__':
    main()
